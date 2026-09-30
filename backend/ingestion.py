"""
Module d'ingestion : extraction de texte brut à partir de fichiers de cours
(PDF ou PPTX) puis découpage en chunks pour l'indexation vectorielle.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import List

from .config import settings


def normalize_typography(text: str) -> str:
    """Remplace les caractères typographiques Unicode par leurs équivalents ASCII.

    Les PDFs de cours utilisent des caractères spécialisés pour l'esthétique,
    ce qui casse la validation du code : "2*n−1" avec un tiret moins (U+2212)
    n'est pas du Python valide. Cette fonction unifie la représentation.

    Substitutions :
    - U+2212 (−) → "-" (tiret moins mathématique)
    - U+2013 (–) → "-" (tiret demi-cadratin)
    - U+2014 (—) → "-" (tiret cadratin)
    - U+2018/2019 ('') → "'" (guillemets simples typographiques)
    - U+201C/201D ("") → '"' (guillemets doubles typographiques)
    - U+00A0 → " " (espace insécable)
    - U+2009 → " " (espace fine)
    - U+2217 (∗) → "*" (astérisque opérateur)
    - U+2026 (…) → "..." (points de suspension)
    """
    # Tirets et tirets moins
    text = text.replace('\u2212', '-')  # U+2212: tiret moins mathématique
    text = text.replace('\u2013', '-')  # U+2013: tiret demi-cadratin (en dash)
    text = text.replace('\u2014', '-')  # U+2014: tiret cadratin (em dash)

    # Guillemets typographiques
    text = text.replace('\u2018', "'")  # U+2018: guillemet simple gauche
    text = text.replace('\u2019', "'")  # U+2019: guillemet simple droit
    text = text.replace('\u201C', '"')  # U+201C: guillemet double gauche
    text = text.replace('\u201D', '"')  # U+201D: guillemet double droit

    # Espaces spécialisées
    text = text.replace('\u00A0', ' ')  # U+00A0: espace insécable
    text = text.replace('\u2009', ' ')  # U+2009: espace fine

    # Astérisque opérateur et points de suspension
    text = text.replace('\u2217', '*')  # U+2217: astérisque opérateur
    text = text.replace('\u2026', '...')  # U+2026: points de suspension

    return text


class UnsupportedFileTypeError(ValueError):
    pass


class EmptyDocumentError(ValueError):
    """Levée quand aucun texte exploitable n'a pu être extrait (ex: PDF scanné sans OCR)."""


@dataclass
class Chunk:
    text: str
    chunk_index: int
    source_page: int | None = None


# --------------------------------------------------------------------------- #
# Extraction
# --------------------------------------------------------------------------- #

def extract_text_from_pdf(path: Path) -> List[tuple[int, str]]:
    """Retourne une liste de (numéro_de_page, texte) pour un PDF."""
    import fitz  # PyMuPDF

    pages: List[tuple[int, str]] = []
    with fitz.open(path) as doc:
        for page_number, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text and text.strip():
                pages.append((page_number, text))
    return pages


def extract_text_from_pptx(path: Path) -> List[tuple[int, str]]:
    """Retourne une liste de (numéro_de_diapo, texte) pour un PPTX."""
    from pptx import Presentation

    slides: List[tuple[int, str]] = []
    prs = Presentation(path)
    for slide_number, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in paragraph.runs)
                    if line.strip():
                        parts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells)
                    if row_text.strip():
                        parts.append(row_text)
        text = "\n".join(parts)
        if text.strip():
            slides.append((slide_number, text))
    return slides


def extract_text(path: Path) -> List[tuple[int, str]]:
    """Détecte le type de fichier et extrait le texte page par page / diapo par diapo."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages = extract_text_from_pdf(path)
    elif suffix in (".pptx",):
        pages = extract_text_from_pptx(path)
    else:
        raise UnsupportedFileTypeError(
            f"Format non supporté : {suffix}. Formats acceptés : .pdf, .pptx"
        )

    if not pages:
        raise EmptyDocumentError(
            "Aucun texte exploitable n'a été trouvé dans ce document. "
            "Le fichier est peut-être un scan sans OCR ou ne contient que des images."
        )

    # Normaliser la typographie (caractères Unicode → ASCII) avant nettoyage.
    # Cela corrige les représentations invalides de code (ex: "2*n−1" avec tiret moins).
    pages = [(page_num, normalize_typography(text)) for page_num, text in pages]

    pages = strip_repeated_boilerplate(pages)
    return pages


# --------------------------------------------------------------------------- #
# Suppression des en-têtes/pieds de page répétés
# --------------------------------------------------------------------------- #

def strip_repeated_boilerplate(
    pages: List[tuple[int, str]], min_repeat_ratio: float = 0.15,
) -> List[tuple[int, str]]:
    """
    Supprime les lignes qui se répètent quasi identiquement sur une grande
    proportion des pages/diapositives (bandeaux de copyright, noms de module,
    numéros de page, etc.) — typiquement "© 2022-2023 – ESPRIT – Module ..."
    présent sur chaque diapositive d'un support de cours.

    Sans ce nettoyage, ce texte répété pollue chaque chunk et peut être capté
    à tort comme "concept clé" par la génération de quiz, en particulier en
    mode mock qui repose sur une simple extraction de mots-clés.

    Deux mécanismes complémentaires :
    1. Toute ligne commençant par "©" est retirée inconditionnellement : un
       bandeau de copyright n'est jamais du contenu pédagogique, quelle que
       soit sa fréquence exacte (utile quand un même document mélange
       plusieurs variantes de bandeau selon les sections, chacune sous le
       seuil de répétition globale — cas réel observé sur des supports
       multi-chapitres).
    2. Toute ligne (hors cas 1) apparaissant sur au moins `min_repeat_ratio`
       des pages est également retirée (numéros de section répétés, etc.).

    Une ligne n'est retirée par le mécanisme 2 que si le document compte au
    moins 3 pages, pour ne pas supprimer de contenu légitime dans un document
    très court.
    """
    def _normalize(line: str) -> str:
        return re.sub(r"\s+", " ", line).strip().lower()

    def _is_copyright_line(norm_line: str) -> bool:
        return norm_line.startswith("©") or norm_line.startswith("(c)")

    # --- Mécanisme 1 : lignes de copyright, retirées sans condition de seuil ---
    stage1_pages: List[tuple[int, str]] = []
    for page_number, text in pages:
        kept_lines = [
            raw_line for raw_line in text.split("\n")
            if not _is_copyright_line(_normalize(raw_line))
        ]
        stage1_pages.append((page_number, "\n".join(kept_lines).strip()))

    if len(pages) < 3:
        return [(pn, t) for pn, t in stage1_pages if t] or pages

    # --- Mécanisme 2 : autres lignes répétées sur >= min_repeat_ratio des pages ---
    line_page_counts: dict[str, set[int]] = {}
    for page_number, text in stage1_pages:
        seen_this_page = set()
        for raw_line in text.split("\n"):
            norm = _normalize(raw_line)
            # On ignore les lignes vides, trop courtes, ou purement numériques
            # (numéros de page) : elles ne posent pas de problème de "faux concept".
            if not norm or len(norm) < 8 or norm.isdigit():
                continue
            seen_this_page.add(norm)
        for norm in seen_this_page:
            line_page_counts.setdefault(norm, set()).add(page_number)

    threshold = max(2, int(len(pages) * min_repeat_ratio))
    boilerplate_lines = {
        norm for norm, page_set in line_page_counts.items() if len(page_set) >= threshold
    }

    cleaned_pages: List[tuple[int, str]] = []
    for page_number, text in stage1_pages:
        kept_lines = [
            raw_line for raw_line in text.split("\n")
            if _normalize(raw_line) not in boilerplate_lines
        ]
        cleaned_text = "\n".join(kept_lines).strip()
        if cleaned_text:
            cleaned_pages.append((page_number, cleaned_text))

    # Filet de sécurité : si le nettoyage a supprimé toutes les pages (cas
    # pathologique), on préfère garder le texte original plutôt que de lever
    # une EmptyDocumentError à tort.
    return cleaned_pages if cleaned_pages else pages


# --------------------------------------------------------------------------- #
# Chunking
# --------------------------------------------------------------------------- #

_WORD_RE = re.compile(r"\S+")
_NEWLINE_TOKEN = "\u23ce"  # marqueur interne représentant un retour à la ligne
_BULLET_START_RE = re.compile(r"^[•\-–▪⮚➢✓*]|^\d+[.)]")
_SENTENCE_END_RE = re.compile(r'[.!?:]["\')]?$')


def _merge_wrapped_lines(text: str) -> str:
    """
    Recolle les lignes qui ne sont que la suite visuelle d'une puce/phrase
    (retour à la ligne dû à la largeur de la diapositive/page PDF), afin que
    chaque ligne du résultat corresponde à une puce ou phrase complète.

    Une ligne est considérée comme une CONTINUATION de la précédente (donc
    fusionnée) si elle ne commence pas par un marqueur de puce/numérotation
    ET que la ligne précédente ne se termine pas par une ponctuation de fin
    de phrase (. ! ? :). Sinon, elle démarre une nouvelle ligne logique.
    """
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    merged: List[str] = []
    for line in lines:
        starts_new_item = bool(_BULLET_START_RE.match(line))
        prev_ends_sentence = bool(merged) and bool(_SENTENCE_END_RE.search(merged[-1]))
        if merged and not starts_new_item and not prev_ends_sentence:
            merged[-1] = f"{merged[-1]} {line}"
        else:
            merged.append(line)
    return "\n".join(merged)


def _approx_token_count(text: str) -> int:
    """Approximation simple : ~1 token ≈ 1 mot (suffisant pour du chunking, pas besoin
    d'un vrai tokenizer BPE ici)."""
    return len(_WORD_RE.findall(text))


def chunk_pages(
    pages: List[tuple[int, str]],
    chunk_size_tokens: int | None = None,
    overlap_tokens: int | None = None,
) -> List[Chunk]:
    """
    Concatène le texte de toutes les pages puis découpe en chunks de taille
    ~chunk_size_tokens avec un chevauchement de ~overlap_tokens, comme spécifié
    dans le cahier des charges (500-1000 tokens, chevauchement 100-200 tokens).

    On conserve la page d'origine dominante de chaque chunk pour la traçabilité.

    Les retours à la ligne originaux sont préservés (via un token interne) à
    travers le découpage mot-à-mot, afin que les phrases/puces individuelles
    restent identifiables dans le texte final — sans cela, la fin d'une puce
    et le titre de la diapositive suivante peuvent se retrouver fusionnés en
    une seule "phrase" incohérente lors d'un traitement en aval (génération
    de quiz, affichage des extraits sources...).
    """
    chunk_size = chunk_size_tokens or settings.CHUNK_SIZE_TOKENS
    overlap = overlap_tokens or settings.CHUNK_OVERLAP_TOKENS
    if overlap >= chunk_size:
        overlap = chunk_size // 4

    # On construit une liste plate de (mot, page_source), en insérant un
    # marqueur dédié à chaque retour à la ligne pour préserver la structure.
    words_with_page: List[tuple[str, int]] = []
    for page_number, text in pages:
        text_with_markers = text.replace("\n", f" {_NEWLINE_TOKEN} ")
        cleaned = re.sub(r"[ \t]+", " ", text_with_markers)
        for word in cleaned.split():
            words_with_page.append((word, page_number))

    if not words_with_page:
        return []

    chunks: List[Chunk] = []
    step = max(chunk_size - overlap, 1)
    idx = 0
    chunk_index = 0
    n = len(words_with_page)
    while idx < n:
        window = words_with_page[idx: idx + chunk_size]
        words = [w for w, _ in window]
        pages_in_window = [p for _, p in window]
        dominant_page = max(set(pages_in_window), key=pages_in_window.count)
        text = " ".join(words).replace(f" {_NEWLINE_TOKEN} ", "\n").replace(_NEWLINE_TOKEN, "\n")
        text = re.sub(r"\n{2,}", "\n", text).strip()
        text = _merge_wrapped_lines(text)
        chunks.append(Chunk(text=text, chunk_index=chunk_index, source_page=dominant_page))
        chunk_index += 1
        idx += step

    return chunks


def process_document(path: Path) -> List[Chunk]:
    """Pipeline complet : extraction + chunking pour un fichier donné."""
    pages = extract_text(path)
    return chunk_pages(pages)
